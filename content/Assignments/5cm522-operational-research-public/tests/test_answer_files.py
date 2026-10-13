"""Published answer files must agree with the code that produces them."""
from __future__ import annotations

import contextlib
import importlib.util
import io
import shutil
import subprocess

import openpyxl
import pytest

from conftest import PYTHON, REPO_ROOT, skip_if_missing

REVISION = REPO_ROOT / "03-revision"

# Sheets written by the code. ("Data and CMA" in Paper A's published workbook also holds
# extra hand-built Excel working, so it is compared separately below.)
CODE_SHEETS = ["Summary", "Seasonal Indices", "Seasonal Variations", "Deseasonalised Trend", "Validation", "Future Forecasts"]

CASES = {
    "a": ("practice_paper_a_code.py", "question_5_forecasting", "practice-paper-a-forecasting-worked.xlsx"),
    "b": ("practice_paper_b_code.py", "question_4_forecasting", "practice-paper-b-forecasting.xlsx"),
}


def load_module(path):
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def assert_same_cells(generated, published, sheets, tol=1e-6):
    g, p = openpyxl.load_workbook(generated), openpyxl.load_workbook(published)
    for sheet in sheets:
        for row in g[sheet].iter_rows():
            for cell in row:
                expected = p[sheet][cell.coordinate].value
                if isinstance(cell.value, (int, float)) and isinstance(expected, (int, float)):
                    assert abs(cell.value - expected) <= tol, (sheet, cell.coordinate, cell.value, expected)
                else:
                    assert cell.value == expected, (sheet, cell.coordinate, cell.value, expected)


@pytest.mark.parametrize("paper", sorted(CASES))
def test_forecasting_workbook_matches_code(paper, tmp_path):
    script, function, workbook = CASES[paper]
    code_path = REVISION / f"practice-paper-{paper}" / script
    skip_if_missing(code_path.read_text(encoding="utf-8"))
    module = load_module(code_path)
    module.WORKBOOK_PATH = tmp_path / "generated.xlsx"
    with contextlib.redirect_stdout(io.StringIO()):
        getattr(module, function)()
    published = REVISION / f"practice-paper-{paper}" / workbook
    assert_same_cells(module.WORKBOOK_PATH, published, CODE_SHEETS)
    # the code-generated columns of "Data and CMA" must match too
    g = openpyxl.load_workbook(module.WORKBOOK_PATH)["Data and CMA"]
    p = openpyxl.load_workbook(published)["Data and CMA"]
    for row in g.iter_rows():
        for cell in row:
            expected = p[cell.coordinate].value
            if isinstance(expected, str) and expected.startswith("="):
                continue  # replaced by an equivalent Excel formula in the worked version
            if expected is None:
                continue  # helper moving averages the worked version deliberately leaves blank
            if isinstance(cell.value, float):
                assert abs(cell.value - expected) <= 1e-6, (cell.coordinate, cell.value, expected)
            else:
                assert cell.value == expected, (cell.coordinate, cell.value, expected)


def test_revision_booklet_is_in_sync_with_practice_papers(tmp_path):
    """The combined booklet is generated from Practice Papers A and B; it must not drift."""
    copy = tmp_path / "03-revision"
    shutil.copytree(REVISION, copy, ignore=shutil.ignore_patterns("*.pdf", "output", "__pycache__"))
    subprocess.run([PYTHON, str(copy / "full-revision-booklet" / "build_full_revision_booklet.py")],
                   check=True, capture_output=True, text=True)
    for name in ("full-revision-booklet.tex", "full-revision-booklet-solutions.tex"):
        fresh = (copy / "full-revision-booklet" / name).read_text(encoding="utf-8").replace("\r\n", "\n")
        committed = (REVISION / "full-revision-booklet" / name).read_text(encoding="utf-8").replace("\r\n", "\n")
        assert fresh == committed, f"{name} is out of date - run build_full_revision_booklet.py"
