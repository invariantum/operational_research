"""Published answer files must agree with the code that produces them."""
from __future__ import annotations

import importlib.util
import json
import re
import subprocess

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import PYTHON, REPO_ROOT, skip_if_missing

TEST_SOLUTIONS = REPO_ROOT / "03-in-class-test" / "solutions"
PART_B = REPO_ROOT / "02-coursework-part-b"


def load_module(path):
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_in_class_test_forecasting_workbook_matches_code(tmp_path):
    code = TEST_SOLUTIONS / "in_class_test_code.py"
    skip_if_missing(code.read_text(encoding="utf-8"))
    module = load_module(code)
    generated = tmp_path / "generated.xlsx"
    module.question_5_results(export_workbook=True, workbook_path=generated)
    g = openpyxl.load_workbook(generated)
    p = openpyxl.load_workbook(TEST_SOLUTIONS / "forecasting-worked-solution.xlsx")
    assert g.sheetnames == p.sheetnames
    for ws in g.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                expected = p[ws.title][cell.coordinate].value
                if isinstance(cell.value, float):
                    assert abs(cell.value - expected) <= 1e-6, (ws.title, cell.coordinate)
                else:
                    assert cell.value == expected, (ws.title, cell.coordinate)


def test_in_class_test_worked_solutions_quote_the_code_results():
    """Key Q5 numbers in the written worked solutions must equal what the code computes."""
    code = TEST_SOLUTIONS / "in_class_test_code.py"
    skip_if_missing(code.read_text(encoding="utf-8"))
    module = load_module(code)
    q5 = module.question_5_results()
    tex = (TEST_SOLUTIONS / "worked-solutions.tex").read_text(encoding="utf-8")
    for key in ("mse_additive", "mse_multiplicative"):
        assert f"{q5['summary'][key]:.4f}" in tex, key
    # The forecast table in the write-up uses rounded intermediate values, so allow 0.01.
    quoted = {int(t): float(v) for t, v in re.findall(r"^(\d+) & Q[1-4] & ([\d.]+) \\\\", tex, re.M)}
    assert len(quoted) == 4, quoted
    for row in q5["future"]:
        assert quoted[row["t"]] == pytest.approx(row["mult_forecast"], abs=0.01), row


def test_create_base_data_reproduces_published_base_data(tmp_path):
    out = tmp_path / "base_data"
    subprocess.run([PYTHON, str(PART_B / "tools" / "create_base_data.py"), "--out_dir", str(out)],
                   check=True, capture_output=True, text=True)
    published = PART_B / "solutions" / "base_data"
    for name in ("zones_base.csv", "stores_base.csv", "zone_store_dist_base.csv"):
        fresh, ref = pd.read_csv(out / name), pd.read_csv(published / name)
        assert list(fresh.columns) == list(ref.columns)
        assert np.allclose(fresh.to_numpy(float), ref.to_numpy(float), rtol=1e-9, atol=1e-9), name
    fresh = json.loads((out / "margins_base.json").read_text())
    ref = json.loads((published / "margins_base.json").read_text())
    assert fresh.keys() == ref.keys()
    for key, value in ref.items():
        if isinstance(value, float):
            assert fresh[key] == pytest.approx(value, rel=1e-12), key
        else:
            assert fresh[key] == value, key
